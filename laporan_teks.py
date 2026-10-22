# -*- coding: utf-8 -*-
"""Fitur 'Buat Laporan dari Teks' — petugas Dit 2.3 menuangkan informasi dari wilayah dalam
bentuk teks bebas; dashboard mengekstrak informasi kunci (tanggal, lokasi, jenis bencana,
angka dampak) lalu menyimpannya ke sesi dan menyediakan unduhan laporan singkat Word."""
import io
import re
from datetime import datetime

import pandas as pd
import streamlit as st
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from data_generator import PROVINSI, WIB
from data_utils import BULAN_PANJANG, fmt_num

NAVY = RGBColor(0x0F, 0x2A, 0x5C)

BULAN_MAP = {b.lower(): i + 1 for i, b in enumerate(BULAN_PANJANG)}
BULAN_MAP.update({"jan": 1, "feb": 2, "mar": 3, "apr": 4, "mei": 5, "jun": 6,
                  "jul": 7, "agu": 8, "sep": 9, "okt": 10, "nov": 11, "des": 12})

PROV_LOWER = {p.lower(): p for p in PROVINSI}
# alias umum
PROV_LOWER.update({
    "sumut": "Sumatera Utara", "sumbar": "Sumatera Barat", "sumsel": "Sumatera Selatan",
    "babel": "Kepulauan Bangka Belitung", "kepri": "Kepulauan Riau",
    "kalbar": "Kalimantan Barat", "kalteng": "Kalimantan Tengah",
    "kalsel": "Kalimantan Selatan", "kaltim": "Kalimantan Timur", "kaltara": "Kalimantan Utara",
})
KAB_LOWER = {}
for p, info in PROVINSI.items():
    for k in info["kab"]:
        KAB_LOWER[k.lower()] = (k, p)

KEYWORDS_TIPE = [
    (["karhutla", "kebakaran hutan", "kebakaran lahan", "titik api", "hotspot", "kebakaran"], "Karhutla"),
    (["banjir", "luapan sungai", "genangan", "rob"], "Banjir"),
    (["gempa", "earthquake", "guncangan", "sesar"], "Gempa"),
    (["longsor", "tanah longsor"], "Bencana Alam Lainnya"),
    (["puting beliung", "angin kencang", "angin puting"], "Bencana Alam Lainnya"),
    (["kekeringan"], "Bencana Alam Lainnya"),
    (["erupsi", "letusan"], "Bencana Alam Lainnya"),
]


# ---------------------------------------------------------------- ekstraksi
def _cari_tanggal(teks):
    t = teks.lower()
    m = re.search(r"\b(\d{1,2})\s+([a-z]{3,9})\s+(\d{4})\b", t)
    if m and m.group(2) in BULAN_MAP:
        try:
            return datetime(int(m.group(3)), BULAN_MAP[m.group(2)], int(m.group(1))).date()
        except ValueError:
            pass
    m = re.search(r"\b(\d{1,2})[-/](\d{1,2})[-/](\d{2,4})\b", t)
    if m:
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if y < 100:
            y += 2000
        try:
            return datetime(y, mo, d).date()
        except ValueError:
            pass
    return None


def _cari_provinsi(teks):
    t = " " + teks.lower() + " "
    hits = []
    for key, nama in PROV_LOWER.items():
        if re.search(r"\b" + re.escape(key) + r"\b", t):
            hits.append(nama)
    seen, out = set(), []
    for n in hits:
        if n not in seen:
            out.append(n)
            seen.add(n)
    return out


def _cari_kabupaten(teks):
    t = " " + teks.lower() + " "
    out = []
    for key, (k, p) in KAB_LOWER.items():
        if re.search(r"\b" + re.escape(key) + r"\b", t):
            out.append((k, p))
    return out


def _cari_jenis(teks):
    t = teks.lower()
    hits = []
    for kws, tipe in KEYWORDS_TIPE:
        for kw in kws:
            if kw in t:
                hits.append(tipe)
                break
    seen, out = set(), []
    for n in hits:
        if n not in seen:
            out.append(n)
            seen.add(n)
    return out


def _angka_dari(pat, teks, group=1, cast=int, flags=re.IGNORECASE):
    m = re.search(pat, teks, flags)
    if not m:
        return None
    s = m.group(group).replace(".", "").replace(",", ".")
    try:
        return cast(float(s))
    except ValueError:
        return None


def _ekstrak_angka(teks):
    out = {}
    v = _angka_dari(r"(\d+(?:[.,]\d+)?)\s*(?:orang|jiwa|warga)?\s*meninggal", teks)
    if v is None:
        v = _angka_dari(r"meninggal\s*(?:dunia)?\s*(\d+)", teks)
    if v is not None:
        out["korban_jiwa"] = v
    v = _angka_dari(r"(\d+(?:[.,]\d+)?)\s*(?:orang|jiwa|warga|kk)\s*(?:mengungsi|pengungsi)", teks)
    if v is None:
        v = _angka_dari(r"pengungsi\s*(?:sebanyak)?\s*(\d+)", teks)
    if v is not None:
        out["pengungsi"] = v
    v = _angka_dari(r"(\d+(?:[.,]\d+)?)\s*(?:unit\s+)?(?:rumah|kk)\s*(?:terdampak|rusak|terendam)", teks)
    if v is None:
        v = _angka_dari(r"rumah\s*(?:terdampak|rusak|terendam)\s*(?:sebanyak)?\s*(\d+)", teks)
    if v is not None:
        out["rumah_terdampak"] = v
    v = _angka_dari(r"(\d+(?:[.,]\d+)?)\s*(?:ha|hektar(?:e)?)\b", teks, cast=float)
    if v is not None:
        out["luas_ha"] = v
    # Luas yang sudah padam / sudah dipadamkan
    v = _angka_dari(r"(\d+(?:[.,]\d+)?)\s*(?:ha|hektar(?:e)?)\s*(?:sudah|telah|berhasil)?\s*(?:di)?padam(?:kan)?",
                    teks, cast=float)
    if v is None:
        v = _angka_dari(r"(?:sudah|telah|berhasil)\s*(?:di)?padam(?:kan)?\D{0,20}(\d+(?:[.,]\d+)?)\s*(?:ha|hektar(?:e)?)",
                        teks, cast=float)
    if v is not None:
        out["luas_padam_ha"] = v
    # Luas yang belum padam (fallback — hitung total-padam kalau total diketahui)
    v = _angka_dari(r"(\d+(?:[.,]\d+)?)\s*(?:ha|hektar(?:e)?)\s*belum\s*(?:di)?padam(?:kan)?",
                    teks, cast=float)
    if v is None:
        v = _angka_dari(r"belum\s*(?:di)?padam(?:kan)?\D{0,20}(\d+(?:[.,]\d+)?)\s*(?:ha|hektar(?:e)?)",
                        teks, cast=float)
    if v is not None:
        out["luas_belum_padam_ha"] = v
    # Jika ada luas_belum_padam tapi tidak ada luas_padam: hitung
    if "luas_belum_padam_ha" in out and "luas_padam_ha" not in out and "luas_ha" in out:
        out["luas_padam_ha"] = max(0.0, out["luas_ha"] - out["luas_belum_padam_ha"])
    v = _angka_dari(r"(?:tinggi\s+(?:genangan|air|banjir)|genangan)\D{0,10}(\d+(?:[.,]\d+)?)\s*cm", teks, cast=float)
    if v is None:
        v = _angka_dari(r"(\d+(?:[.,]\d+)?)\s*cm", teks, cast=float)
    if v is not None:
        out["tinggi_genangan_cm"] = v
    v = _angka_dari(r"(?:magnitudo|skala\s+richter|sr|m\s*=?)\s*([0-9]+(?:[.,][0-9]+)?)", teks, cast=float)
    if v is not None:
        out["magnitudo"] = v
    v = _angka_dari(r"(?:kedalaman)\D{0,10}(\d+(?:[.,]\d+)?)\s*km", teks, cast=float)
    if v is not None:
        out["kedalaman_km"] = v
    v = _angka_dari(r"(\d+(?:[.,]\d+)?)\s*(?:kasus\s+)?ispa", teks)
    if v is not None:
        out["ispa"] = v
    v = _angka_dari(r"(\d+(?:[.,]\d+)?)\s*(?:titik\s+api|hotspot)", teks)
    if v is not None:
        out["hotspot"] = v
    return out


def analisis_teks(teks):
    """Pecah teks laporan menjadi struktur: tanggal, provinsi, kabupaten, jenis, angka, teks_asli."""
    teks = (teks or "").strip()
    if not teks:
        return None
    tgl = _cari_tanggal(teks)
    provs = _cari_provinsi(teks)
    kabs = _cari_kabupaten(teks)
    jenis = _cari_jenis(teks)
    if not provs and kabs:
        provs = sorted({p for _, p in kabs})
    angka = _ekstrak_angka(teks)
    return {
        "tanggal": tgl,
        "provinsi": provs,
        "kabupaten": [k for k, _ in kabs],
        "jenis_bencana": jenis,
        "angka": angka,
        "teks_asli": teks,
    }


# ---------------------------------------------------------------- docx
def build_laporan_teks_docx(rec):
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    st_ = doc.styles["Normal"]
    st_.font.name = "Calibri"
    st_.font.size = Pt(11)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("BADAN INTELIJEN NEGARA REPUBLIK INDONESIA")
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = NAVY
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Deputi 2 · Direktorat 2.3")
    r.font.size = Pt(11)
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run("LAPORAN KEJADIAN DARI WILAYAH")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = NAVY

    tgl = rec["ekstraksi"].get("tanggal")
    tgl_s = f"{tgl.day} {BULAN_PANJANG[tgl.month - 1]} {tgl.year}" if tgl else "-"
    meta = [
        ("Judul", rec.get("judul", "-")),
        ("Tanggal kejadian", tgl_s),
        ("Provinsi", ", ".join(rec["ekstraksi"]["provinsi"]) or "-"),
        ("Kabupaten/Kota", ", ".join(rec["ekstraksi"]["kabupaten"]) or "-"),
        ("Jenis bencana", ", ".join(rec["ekstraksi"]["jenis_bencana"]) or "-"),
        ("Sumber", rec.get("sumber", "-")),
        ("Petugas Dit 2.3", rec.get("petugas", "-")),
        ("Dibuat", rec.get("dibuat", "-")),
    ]
    t = doc.add_table(rows=0, cols=2)
    for k, v in meta:
        c = t.add_row().cells
        c[0].text = k
        c[1].text = f": {v}"
        c[0].width, c[1].width = Cm(4.5), Cm(11.5)

    doc.add_heading("Ringkasan Dampak (hasil ekstraksi otomatis)", level=2)
    angka = rec["ekstraksi"]["angka"]
    label = {
        "korban_jiwa": "Korban jiwa (orang)", "pengungsi": "Pengungsi (orang)",
        "rumah_terdampak": "Rumah terdampak (unit)", "luas_ha": "Luas terbakar (ha)",
        "luas_padam_ha": "Luas sudah padam (ha)", "luas_belum_padam_ha": "Luas belum padam (ha)",
        "tinggi_genangan_cm": "Tinggi genangan (cm)", "magnitudo": "Magnitudo (SR)",
        "kedalaman_km": "Kedalaman (km)", "ispa": "Kasus ISPA", "hotspot": "Titik api / hotspot",
    }
    if angka:
        tb = doc.add_table(rows=0, cols=2)
        tb.style = "Light Grid Accent 1"
        for k, v in angka.items():
            c = tb.add_row().cells
            c[0].text = label.get(k, k)
            c[1].text = fmt_num(v, 1 if isinstance(v, float) and not float(v).is_integer() else 0)
    else:
        doc.add_paragraph("Tidak terdeteksi angka dampak dalam teks.")

    doc.add_heading("Narasi Petugas Wilayah", level=2)
    doc.add_paragraph(rec["ekstraksi"]["teks_asli"])

    p = doc.add_paragraph()
    r = p.add_run("Laporan ini perlu diverifikasi sebelum dieskalasi ke pimpinan.")
    r.italic = True
    r.font.size = Pt(9)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- UI
def _state():
    if "manual_reports" not in st.session_state:
        st.session_state["manual_reports"] = []
    return st.session_state["manual_reports"]


def _tipe_ke_df_row(ek, petugas):
    """Konversi hasil ekstraksi menjadi 1 baris data kejadian (per jenis & provinsi)."""
    rows = []
    tgl = ek.get("tanggal") or datetime.now(WIB).date()
    jenis_list = ek["jenis_bencana"] or ["Bencana Alam Lainnya"]
    prov_list = ek["provinsi"] or ["-"]
    kab = ek["kabupaten"][0] if ek["kabupaten"] else "-"
    a = ek["angka"]
    for prov in prov_list:
        info = PROVINSI.get(prov, {"lat": -0.5, "lon": 108.5})
        for j in jenis_list:
            rows.append({
                "id_kejadian": f"T-{int(datetime.now(WIB).timestamp())}-{len(rows)}",
                "tanggal": pd.Timestamp(tgl),
                "provinsi": prov,
                "kabupaten": kab,
                "lat": info["lat"],
                "lon": info["lon"],
                "tipe_bencana": j,
                "luas_ha": a.get("luas_ha"),
                "luas_padam_ha": a.get("luas_padam_ha"),
                "magnitudo": a.get("magnitudo"),
                "kedalaman_km": a.get("kedalaman_km"),
                "tinggi_genangan_cm": a.get("tinggi_genangan_cm"),
                "rumah_terdampak": a.get("rumah_terdampak"),
                "korban_jiwa": int(a.get("korban_jiwa", 0) or 0),
                "pengungsi": int(a.get("pengungsi", 0) or 0),
                "ispa": int(a.get("ispa", 0) or 0),
                "penyebab": "-",
                "status_verifikasi": "Perlu Verifikasi",
                "sumber": f"Laporan teks ({petugas})",
            })
    return rows


def render_tab_buat_laporan(user, admin, base_df=None):
    st.markdown("Tulis atau tempel laporan dari petugas wilayah. Dashboard akan **mengekstrak** "
                "tanggal, lokasi, jenis bencana, dan angka dampak secara otomatis.")

    with st.form("laporan_teks_form", clear_on_submit=False):
        c1, c2 = st.columns(2)
        judul = c1.text_input("Judul singkat laporan", placeholder="Mis. Karhutla di Pelalawan 18 Okt 2026")
        sumber = c2.text_input("Sumber informasi", placeholder="Mis. Posko BPBD Riau / Danramil / Camat")
        teks = st.text_area("Teks laporan dari wilayah",
                            placeholder=("Contoh: Pada 18 Oktober 2026 terjadi karhutla di Kabupaten Pelalawan, "
                                         "Provinsi Riau. Luas lahan terbakar sekitar 90 ha, 80 ha sudah "
                                         "dipadamkan dan 10 ha belum padam. 45 warga mengungsi, 18 kasus ISPA "
                                         "tercatat. Terpantau 12 titik api di sekitar lokasi."),
                            height=180)
        c1, c2, c3 = st.columns([1, 1, 1])
        pratinjau = c1.form_submit_button("🔍 Pratinjau Ekstraksi")
        simpan = c2.form_submit_button("💾 Simpan Laporan", type="primary")
        tambahkan_ke_data = c3.checkbox("Tambahkan juga ke data dashboard (sesi)",
                                        value=False, help="Baris baru akan muncul di tabel/peta/rekap. "
                                                          "Hanya aktif selama sesi berjalan.")

    if pratinjau or simpan:
        ek = analisis_teks(teks)
        if ek is None:
            st.warning("Teks laporan masih kosong.")
            return
        _tampil_ekstraksi(ek)

        if simpan:
            rec = {
                "id": f"LP-{int(datetime.now(WIB).timestamp())}",
                "judul": judul.strip() or "(tanpa judul)",
                "sumber": sumber.strip() or "-",
                "petugas": f"{user['name']} ({user['role']})",
                "dibuat": datetime.now(WIB).strftime("%Y-%m-%d %H:%M WIB"),
                "ekstraksi": ek,
            }
            _state().insert(0, rec)
            st.success(f"Laporan **{rec['judul']}** tersimpan ke sesi.")
            if tambahkan_ke_data:
                rows = _tipe_ke_df_row(ek, user["name"])
                extra = pd.DataFrame(rows)
                # Base: pakai up_inc existing; kalau belum ada, pakai base_df dari halaman
                # (biasanya data demo yang sedang aktif) agar data lama tidak hilang.
                base = st.session_state.get("up_inc")
                if base is None and base_df is not None and not base_df.empty:
                    base = base_df.copy()
                merged = pd.concat([base, extra], ignore_index=True) if base is not None else extra
                st.session_state["up_inc"] = merged
                st.session_state["src"] = "upload"
                st.info(f"{len(rows)} baris kejadian ditambahkan ke dataset dashboard "
                        f"(total sekarang: {len(merged)} baris).")

    # Indikator mode dataset + tombol reset
    if st.session_state.get("src") == "upload" and st.session_state.get("up_inc") is not None:
        n_rows = len(st.session_state["up_inc"])
        colA, colB = st.columns([3, 1])
        colA.info(f"ℹ️ Dataset dashboard saat ini mode **upload** ({n_rows} baris). "
                  "Berisi data demo + baris laporan teks yang pernah ditambahkan.")
        if colB.button("🔄 Reset ke data demo", key="reset_demo",
                       help="Hapus tambahan laporan teks dari dataset dashboard dan kembali ke data demo."):
            st.session_state.pop("up_inc", None)
            st.session_state.pop("up_hs", None)
            st.session_state.pop("sig_inc", None)
            st.session_state.pop("sig_hs", None)
            st.session_state["src"] = "demo"
            st.success("Dataset dikembalikan ke data demo.")
            st.rerun()

    # Daftar laporan tersimpan di sesi
    daftar = _state()
    st.divider()
    st.markdown(f"**Laporan tersimpan dalam sesi: {len(daftar)}**")
    if not daftar:
        st.caption("Belum ada laporan. Simpan laporan untuk mengisi daftar ini.")
        return
    for i, rec in enumerate(daftar):
        with st.expander(f"{rec['judul']}  ·  {rec['dibuat']}", expanded=(i == 0)):
            ek = rec["ekstraksi"]
            cA, cB = st.columns(2)
            cA.markdown(f"**Petugas:** {rec['petugas']}  \n**Sumber:** {rec['sumber']}")
            tgl_s = (f"{ek['tanggal'].day} {BULAN_PANJANG[ek['tanggal'].month - 1]} {ek['tanggal'].year}"
                     if ek.get("tanggal") else "-")
            cB.markdown(f"**Tanggal kejadian:** {tgl_s}  \n"
                        f"**Jenis bencana:** {', '.join(ek['jenis_bencana']) or '-'}  \n"
                        f"**Provinsi:** {', '.join(ek['provinsi']) or '-'}")
            if ek["angka"]:
                st.markdown("**Angka dampak terdeteksi**")
                st.dataframe(pd.DataFrame([ek["angka"]]).T.rename(columns={0: "Nilai"}),
                             use_container_width=True)
            st.caption("Teks laporan:")
            st.write(ek["teks_asli"])
            c1, c2 = st.columns([1, 1])
            c1.download_button("⬇️ Unduh laporan singkat (.docx)", build_laporan_teks_docx(rec),
                               f"laporan_{rec['id']}.docx",
                               "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                               key=f"dl_{rec['id']}")
            if admin and c2.button("🗑️ Hapus laporan ini", key=f"del_{rec['id']}"):
                _state().remove(rec)
                st.rerun()


def _tampil_ekstraksi(ek):
    with st.container(border=True):
        st.markdown("**Hasil ekstraksi otomatis**")
        tgl_s = (f"{ek['tanggal'].day} {BULAN_PANJANG[ek['tanggal'].month - 1]} {ek['tanggal'].year}"
                 if ek.get("tanggal") else "_tidak terdeteksi_")
        c1, c2, c3 = st.columns(3)
        c1.markdown(f"**Tanggal kejadian**  \n{tgl_s}")
        c2.markdown(f"**Jenis bencana**  \n{', '.join(ek['jenis_bencana']) or '_tidak terdeteksi_'}")
        c3.markdown(f"**Provinsi**  \n{', '.join(ek['provinsi']) or '_tidak terdeteksi_'}")
        if ek["kabupaten"]:
            st.markdown(f"**Kabupaten/Kota terdeteksi:** {', '.join(ek['kabupaten'])}")
        if ek["angka"]:
            st.markdown("**Angka dampak terdeteksi:**")
            st.dataframe(pd.DataFrame([ek["angka"]]).T.rename(columns={0: "Nilai"}),
                         use_container_width=True)
        else:
            st.caption("Tidak terdeteksi angka dampak. Periksa kembali ejaan satuan "
                       "(ha, cm, km, SR) dan kata kunci (meninggal, mengungsi, rumah, ISPA, titik api).")
